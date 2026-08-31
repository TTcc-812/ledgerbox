from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

AMOUNT_RE = re.compile(r"-?[\d,.]+")


@dataclass
class Txn:
    id: str
    time: str
    platform: str
    direction: str
    amount: float
    counterparty: str
    description: str
    method: str
    status: str
    category: str
    raw_type: str

    def to_dict(self) -> dict:
        return asdict(self)


def _decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030", "gbk", "utf-8"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _money(value: object) -> float:
    if value is None:
        return 0.0
    s = str(value).replace("¥", "").replace(",", "").replace("￥", "").strip()
    m = AMOUNT_RE.search(s)
    if not m:
        return 0.0
    try:
        return abs(float(m.group().replace(",", "")))
    except ValueError:
        return 0.0


def _time(value: object) -> str:
    s = str(value or "").strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d",
    ):
        try:
            chunk = s[:19] if len(s) >= 19 else s
            return datetime.strptime(chunk, fmt).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    return s


def _id(*parts: str) -> str:
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()[:16]


def _rows_from_file(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    text = _decode_bytes(path.read_bytes())
    return [list(r) for r in csv.reader(io.StringIO(text))]


def detect_platform(path: Path, rows: list[list[str]]) -> str:
    blob = path.name + " " + " ".join(" ".join(r[:6]) for r in rows[:30])
    if "微信支付" in blob or "微信昵称" in blob:
        return "wechat"
    if "支付宝" in blob or "电子客户回单" in blob:
        return "alipay"
    if "招商银行" in blob or "招行" in blob:
        return "cmb"
    head = ",".join(rows[0][:8]) if rows else ""
    if "交易对方" in head and "金额" in head:
        return "wechat"
    if "收/付款方式" in head or "交易分类" in head:
        return "alipay"
    return "cmb"


def _find_header(rows: list[list[str]], keys: tuple[str, ...]) -> int:
    for i, row in enumerate(rows):
        joined = "".join(row)
        if all(k in joined for k in keys):
            return i
    return -1


def parse_file(path: Path) -> list[Txn]:
    rows = _rows_from_file(path)
    platform = detect_platform(path, rows)
    if platform == "wechat":
        return _parse_wechat(rows)
    if platform == "alipay":
        return _parse_alipay(rows)
    return _parse_cmb(rows)


def _colmap(header: list[str]) -> dict[str, int]:
    return {h.strip(): i for i, h in enumerate(header)}


def _cell(row: list[str], idx: int) -> str:
    return row[idx].strip() if 0 <= idx < len(row) else ""


def _parse_wechat(rows: list[list[str]]) -> list[Txn]:
    hi = _find_header(rows, ("交易时间", "金额"))
    if hi < 0:
        return []
    m = _colmap(rows[hi])
    out: list[Txn] = []
    for row in rows[hi + 1 :]:
        if len(row) < 5:
            continue
        t = _time(_cell(row, m.get("交易时间", 0)))
        if not t:
            continue
        direction = _cell(row, m.get("收/支", 4))
        amount = _money(_cell(row, m.get("金额(元)", m.get("金额", 5))))
        counter = _cell(row, m.get("交易对方", 2))
        desc = _cell(row, m.get("商品", 3))
        method = _cell(row, m.get("支付方式", 6))
        status = _cell(row, m.get("当前状态", 7))
        oid = _cell(row, m.get("交易单号", 8)) or _id(t, counter, str(amount), desc)
        raw_type = _cell(row, m.get("交易类型", 1))
        if direction not in ("支出", "收入"):
            direction = "不计"
        out.append(
            Txn(oid.replace("\t", ""), t, "wechat", direction, amount, counter, desc, method, status, "", raw_type)
        )
    return out


def _parse_alipay(rows: list[list[str]]) -> list[Txn]:
    hi = _find_header(rows, ("交易时间", "金额"))
    if hi < 0:
        return []
    m = _colmap(rows[hi])
    out: list[Txn] = []
    for row in rows[hi + 1 :]:
        if len(row) < 6:
            continue
        t = _time(_cell(row, m.get("交易时间", 0)))
        if not t:
            continue
        direction = _cell(row, m.get("收/支", 5))
        amount = _money(_cell(row, m.get("金额", 6)))
        counter = _cell(row, m.get("交易对方", 2))
        desc = _cell(row, m.get("商品说明", 4))
        method = _cell(row, m.get("收/付款方式", 7))
        status = _cell(row, m.get("交易状态", 8))
        oid = _cell(row, m.get("交易订单号", 9)) or _id(t, counter, str(amount), desc)
        raw_type = _cell(row, m.get("交易分类", 1))
        if direction not in ("支出", "收入"):
            direction = "不计"
        out.append(
            Txn(oid.replace("\t", ""), t, "alipay", direction, amount, counter, desc, method, status, "", raw_type)
        )
    return out


def _parse_cmb(rows: list[list[str]]) -> list[Txn]:
    hi = -1
    for i, row in enumerate(rows):
        joined = "".join(row)
        if "交易日" in joined or "交易日期" in joined or "记账日期" in joined:
            hi = i
            break
    if hi < 0:
        return []
    m = _colmap(rows[hi])

    def pick(*names: str) -> int:
        for n in names:
            if n in m:
                return m[n]
        return -1

    i_date = pick("交易日", "交易日期", "记账日期", "日期")
    i_time = pick("交易时间", "时间")
    i_out = pick("支出", "借方发生额", "借方金额")
    i_in = pick("收入", "贷方发生额", "贷方金额")
    i_amt = pick("交易金额", "金额")
    i_type = pick("交易类型", "摘要")
    i_note = pick("交易备注", "备注", "摘要")
    i_cp = pick("对方户名", "对方账号", "对手信息")
    out: list[Txn] = []
    for row in rows[hi + 1 :]:
        day = _cell(row, i_date)
        if not day:
            continue
        clock = _cell(row, i_time) if i_time >= 0 else "00:00:00"
        t = _time(f"{day} {clock}".strip())
        spent = _money(_cell(row, i_out)) if i_out >= 0 else 0.0
        gained = _money(_cell(row, i_in)) if i_in >= 0 else 0.0
        if spent:
            direction, amount = "支出", spent
        elif gained:
            direction, amount = "收入", gained
        else:
            signed = str(_cell(row, i_amt))
            amount = _money(signed)
            direction = "支出" if "-" in signed else "收入"
        counter = _cell(row, i_cp)
        desc = _cell(row, i_note) or _cell(row, i_type)
        oid = _id("cmb", t, str(amount), counter, desc)
        out.append(
            Txn(oid, t, "cmb", direction, amount, counter, desc, "招商银行", "成功", "", _cell(row, i_type))
        )
    return out
